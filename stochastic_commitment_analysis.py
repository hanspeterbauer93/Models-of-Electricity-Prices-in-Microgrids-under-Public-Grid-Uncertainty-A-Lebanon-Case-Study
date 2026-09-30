"""
What the microgrid's own build does to the leader's stochastic commitment.

:mod:`stochastic_analysis` sweeps the leader's BELIEF and reports the
stochastic commitment -- all six EDL outcomes at their real probabilities, one
price committed before the day starts -- as a single reference point on each
figure. This module holds the belief fixed at that real one and moves the
MICROGRID instead. Two comparisons, each two builds:

    fleet   the genset alone (CAP["PV"] forced to 0) against the same genset
            plus PV_KWP of PV. Both keep the configured load shift.
    shift   the configured load shift (10% of each hour's demand) against a
            raised one (50%). Both carry PV_KWP of PV, since the point of the
            comparison is how much a flexible village is worth to a microgrid
            that owns a solar fleet it must use when the sun allows.

The fleet is this module's own -- :data:`DG_KW` and :data:`PV_KWP`, patched onto
the config for the duration of each solve and restored afterwards. The genset
matches the case study in :mod:`config`; the PV does not, since config carries
CAP["PV"] = 0. Patching rather than editing config is what lets the fleet be
resized here (--dg / --pv) without moving any other module's results.

and for each build solves the two runs the thesis reads against each other:

    no supply   P(none) = 1: EDL is off all day, so the microgrid faces the
                whole village load. The MGO's best case for volume, and the
                pessimistic end of the belief range.
    commitment  all six outcomes at their probabilities from
                :data:`scenarios.SCENARIOS`. The commitment the MGO actually
                makes, since it must post one price without knowing which
                outcome arrives.

THE SPAN
The quantity plotted is the distance between those two runs,

    span = commitment - no supply       [$/kWh, or $/day]
    span% = 100 * span / (no supply)    [% of the EDL-off day]

for the leader price and for the leader profit, drawn as one horizontal range
per (season, build). The span is what EDL's expected presence is worth to the
leader relative to a day it stays off: the price span says how far EDL's
arrival moves the tariff the village pays, the profit span how far it moves the
MGO's daily earnings. The percentage is the same number against the end the
line starts from, which is what makes a price span and a profit span -- three
orders of magnitude apart in units -- comparable at all.

The two builds of a comparison are drawn on the same axis, so the effect of PV
(or of flexibility) is the DIFFERENCE BETWEEN TWO SPANS rather than a level
shift that could be read off either run alone. Their difference is reported in
percentage POINTS, since the spans are already percentages.

Profit spans are expected to run leftward (negative): EDL delivering carves
demand out of the microgrid's market, so the commitment earns less than the
no-supply run. That is a statement about direction, not magnitude, and the
figures report the magnitude.

WHAT TO LOOK FOR
PV enters on the COST side only -- it is another way to serve the same kWh, at
c_OV["PV"] instead of the diesel marginal cost. The load shift enters on the
VOLUME side: it lets the village move up to shift_frac of each hour's demand,
which under an EDL window that covers only part of the day means moving load
INTO the hours EDL serves and out of the hours the microgrid would have sold
into. And the residual-based demand response (see :mod:`model.bilevel_mpec`)
makes EDL availability scale the microgrid's residual demand rather than shift
it.

If all three hold, both knobs should move the profit spans far more than the
price spans, and the load shift should bite hardest in the season where the EDL
window and the demand peak line up worst. A price span that either knob changes
materially is the interesting finding, because it would mean the microgrid's
own build, not the market it faces, is setting the tariff.

Within a comparison the two builds differ in ONE config entry and nothing else;
the demand data, the scenario probabilities and the price menu are shared, so
any difference between the rows is that entry.

Outputs, under :data:`config.OUTPUT_DIR`:

    stochastic_commitment_analysis.xlsx
            both spans for every build, the build-to-build difference,
            and every solved run in full
    stochastic_commitment_price_span_<comparison>.png
            price span per season and build
    stochastic_commitment_profit_span_<comparison>.png
            profit span per season and build

Usage:
    python stochastic_commitment_analysis.py                     # both comparisons
    python stochastic_commitment_analysis.py --comparisons shift
    python stochastic_commitment_analysis.py --seasons winter
    python stochastic_commitment_analysis.py --dg 400 --pv 400 --shift 0.4
    python stochastic_commitment_analysis.py --levels 129
"""

import argparse
from contextlib import contextmanager
from dataclasses import dataclass

import matplotlib
matplotlib.use("Agg")             # headless / file output
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import config
# The two runs are exactly the ones :mod:`stochastic_analysis` already draws as
# reference lines, imported rather than re-derived so the numbers here and the
# numbers on those figures are the same runs.
from stochastic_analysis import FIELDS, commitment_run, no_supply_reference

_TOL = 1e-6

# The microgrid's own generation, for every build here.           [ASSUMPTION]
#
# Deliberately NOT config.CONFIG["CAP"]: this module sizes the fleet itself and
# patches the config for the duration of each solve, so the fleet can be resized
# here without moving the case study every other module is built on. The genset
# matches config's 400 kW; the PV does not, since config carries CAP["PV"] = 0.
#
# Matched to each other so the fleet comparison is "the same nameplate, twice
# over": firm diesel against PV that only delivers into the capacity factor.
# Overridable with --dg / --pv.
DG_KW: float = 400.0
PV_KWP: float = 400.0

# Solver budget for this module's solves.                         [ASSUMPTION]
# The raised-shift builds are a much harder MILP than anything else here -- a
# wider shift band frees 24 more continuous variables against the same
# complementarity binaries -- and at config's 300 s they abort at the time
# limit, which drops the build and leaves the shift comparison with one row.
# The gap is config's; only the clock is longer. A build that still aborts is
# reported rather than silently missing, so raise this again if one does.
TIME_LIMIT_S: float = 1800.0

# Default raised load shift of the shift comparison [-].          [ASSUMPTION]
# Five times the configured 10%: not a plausible policy so much as a bound on
# what village flexibility could ever be worth here. Overridable with --shift.
SHIFT_HIGH: float = 0.50

# Two Okabe-Ito hues per comparison, chosen so the fleet figures and the shift
# figures cannot be confused with each other at a glance -- and kept clear of
# :data:`plot.SCENARIO_COLOR`, because these rows are BUILDS, not EDL outcomes.
_INK = {"dg": "#D55E00",        # vermillion
        "dg_pv": "#009E73",     # green
        "shift_std": "#0072B2",  # blue
        "shift_high": "#E69F00"}  # orange

# The two runs each build is solved for, in the order they are written.
RUNS = (("reference", "no supply (P = 1)"),
        ("commitment", "stochastic commitment (real P)"))


# --------------------------------------------------------------------------
# The builds under comparison
# --------------------------------------------------------------------------
@dataclass(frozen=True)
class Build:
    """One configuration of the microgrid, as the config entries to force."""
    key: str
    label: str
    dg_kw: float
    pv_kwp: float
    shift_frac: float

    @property
    def color(self) -> str:
        return _INK.get(self.key, "#444444")

    @property
    def stamp(self) -> tuple:
        """What actually goes into the model, so two builds that differ only in
        name are solved once."""
        return (round(self.dg_kw, 6), round(self.pv_kwp, 6),
                round(self.shift_frac, 6))


@dataclass(frozen=True)
class Comparison:
    """Two builds differing in one config entry, and what that entry is."""
    key: str            # file stem
    what: str           # for the figure title
    fixed: str          # what is held equal, for the caption
    builds: tuple[Build, ...]


def _dg_label(dg_kw: float) -> str:
    return f"{dg_kw:.0f} kW DG"


def _fleet_label(dg_kw: float, pv_kwp: float) -> str:
    return f"{_dg_label(dg_kw)} + {pv_kwp:.0f} kWp PV"


def _shift_label(frac: float) -> str:
    std = " (standard)" if abs(frac - config.CONFIG["shift_frac"]) <= _TOL else ""
    return f"{100 * frac:.0f}% load shift{std}"


def fleet_comparison(dg_kw: float = DG_KW,
                     pv_kwp: float = PV_KWP) -> Comparison:
    """Genset alone against genset plus PV, at the configured load shift."""
    shift = config.CONFIG["shift_frac"]
    return Comparison(
        key="fleet",
        what="with and without PV",
        fixed=f"Both builds carry the configured {100 * shift:.0f}% load "
              f"shift; they differ in CAP['PV'] and in nothing else.",
        builds=(
            Build("dg", _dg_label(dg_kw), dg_kw, 0.0, shift),
            Build("dg_pv", _fleet_label(dg_kw, pv_kwp), dg_kw, pv_kwp, shift),
        ),
    )


def shift_comparison(dg_kw: float = DG_KW, pv_kwp: float = PV_KWP,
                     shift_high: float = SHIFT_HIGH) -> Comparison:
    """The configured load shift against a raised one, both on the PV fleet."""
    std = config.CONFIG["shift_frac"]
    return Comparison(
        key="shift",
        what=f"at {100 * std:.0f}% and {100 * shift_high:.0f}% load shift",
        fixed=f"Both builds are {_fleet_label(dg_kw, pv_kwp)}; they differ in "
              f"shift_frac and in nothing else.",
        builds=(
            Build("shift_std", _shift_label(std), dg_kw, pv_kwp, std),
            Build("shift_high", _shift_label(shift_high), dg_kw, pv_kwp,
                  shift_high),
        ),
    )


COMPARISONS = {"fleet": fleet_comparison, "shift": shift_comparison}


# --------------------------------------------------------------------------
# Solving
# --------------------------------------------------------------------------
@contextmanager
def build_settings(build: Build):
    """
    Run the body with the config forced to ``build``'s generation capacities
    and load shift.

    :func:`run.build_instance` reads ``config.CONFIG`` at call time, and the
    imported solve functions build their own instances, so the settings are
    patched on the config rather than on a data dict. ``CAP["DG"]`` feeds the
    diesel cap, ``CAP["PV"]`` feeds QPV[t] = kwp * phi_t, both feed the PV and
    diesel terms of the objective and the big-M on qMGO; ``shift_frac`` feeds
    the shift bounds and the big-Ms on qsupp and the shift slack -- so this is
    a real change of build, not a relabelled run. All three are restored on the
    way out, including if the solve raises.
    """
    cfg = config.CONFIG
    old = (cfg["CAP"]["DG"], cfg["CAP"]["PV"], cfg["shift_frac"])
    cfg["CAP"]["DG"] = build.dg_kw
    cfg["CAP"]["PV"] = build.pv_kwp
    cfg["shift_frac"] = build.shift_frac
    try:
        yield
    finally:
        cfg["CAP"]["DG"], cfg["CAP"]["PV"], cfg["shift_frac"] = old


def solve_build(season: str, build: Build, opts: dict, levels: int | None,
                cache: dict | None = None) -> dict | None:
    """
    Both runs of one season under one build, or ``None`` if either failed.

    A build with only one of the two runs has no span, which is the whole
    quantity this module reports, so it is dropped rather than half-drawn.

    ``cache`` is keyed by what actually reaches the model, so a build shared by
    two comparisons (the PV fleet is both ``dg_pv`` and ``shift_std``) is
    solved once and read twice.
    """
    ckey = (season, build.stamp)
    if cache is not None and ckey in cache:
        print(f"[{season}] {build.label}: same model as a build already "
              f"solved — reusing it")
        runs = cache[ckey]
    else:
        print(f"[{season}] {build.label}")
        with build_settings(build):
            runs = (no_supply_reference(season, opts, levels),
                    commitment_run(season, opts, levels))
        if cache is not None:
            cache[ckey] = runs

    ref, com = runs
    if ref is None or com is None:
        missing = ", ".join(n for n, r in (("no supply", ref),
                                              ("stochastic commitment", com))
                            if r is None)
        print(f"  {build.label}: {missing} did not solve — no span, dropped")
        return None

    return {"season": season, "build": build.key, "label": build.label,
            "dg_kw": build.dg_kw, "pv_kwp": build.pv_kwp,
            "shift_frac": build.shift_frac, "reference": ref, "commitment": com}


def span(blob: dict, key: str) -> float | None:
    """``commitment - no supply`` of one field, or ``None`` if either end has no
    value (a degenerate run sells nothing, so it carries no price)."""
    lo, hi = blob["reference"][key], blob["commitment"][key]
    if lo is None or hi is None:
        return None
    d = hi - lo
    # Both runs on the same price-menu level differ by float noise only, which
    # would otherwise print as a signed "-0.0%".
    return 0.0 if abs(d) <= _TOL else d


def _pct(value_: float | None, base: float | None) -> float | None:
    """``value_`` as a percentage of ``base``, or ``None`` where that is not a
    number (no value at one end, or an EDL-off day that earned nothing)."""
    if value_ is None or base is None or abs(base) <= _TOL:
        return None
    return 100.0 * value_ / base


def summarize(blob: dict) -> dict:
    """Both spans of one (season, build): absolute, as a share of the
    no-supply end, and the two runs they are measured between."""
    ref, com = blob["reference"], blob["commitment"]
    P_ref = config.CONFIG["P_ref"]
    price_span, profit_span = span(blob, "p"), span(blob, "profit")

    return {
        "season": blob["season"],
        "build": blob["build"],
        "label": blob["label"],
        "dg_kw": blob["dg_kw"],
        "pv_kwp": blob["pv_kwp"],
        "shift_pct": 100.0 * blob["shift_frac"],
        "p_none": ref["p"],
        "p_commit": com["p"],
        "price_span": price_span,
        "price_span_x": None if price_span is None else price_span / P_ref,
        "price_span_pct": _pct(price_span, ref["p"]),
        "profit_none": ref["profit"],
        "profit_commit": com["profit"],
        "profit_span": profit_span,
        "profit_span_pct": _pct(profit_span, ref["profit"]),
        "qMGO_commit": com["qMGO"],
        "qPG_commit": com["qPG"],
        "qsupp_commit": com["qsupp"],
        "qcri_commit": com["qcri"],
    }


def contrast(rows: list[dict], cmp_: Comparison) -> dict | None:
    """
    What the changed config entry did, for one season: the second build of the
    comparison minus the first.

    Returns ``None`` unless both builds are present -- with one of them missing
    there is nothing to difference. Differences of the percentage spans are in
    percentage POINTS, since both terms are already percentages.
    """
    by_key = {r["build"]: r for r in rows}
    first, second = (b.key for b in cmp_.builds)
    if first not in by_key or second not in by_key:
        return None
    base, alt = by_key[first], by_key[second]

    def diff(key):
        a, b = base[key], alt[key]
        if a is None or b is None:
            return None
        d = b - a
        # Two builds that landed on the same menu level differ by float noise
        # only (~1e-15 on a price of ~0.6, against a menu step of ~0.013).
        # Reporting that as a number would invite reading a sign into it.
        return 0.0 if abs(d) <= _TOL else d

    return {
        "comparison": cmp_.key,
        "season": base["season"],
        "against": f"{alt['label']} minus {base['label']}",
        "d_p_none": diff("p_none"),
        "d_p_commit": diff("p_commit"),
        "d_price_span": diff("price_span"),
        "d_price_span_pp": diff("price_span_pct"),
        "d_profit_none": diff("profit_none"),
        "d_profit_commit": diff("profit_commit"),
        "d_profit_span": diff("profit_span"),
        "d_profit_span_pp": diff("profit_span_pct"),
    }


# --------------------------------------------------------------------------
# Summary
# --------------------------------------------------------------------------
def _fmt(v, spec: str, width: int) -> str:
    """A number, or a dash of the same width where there is none."""
    return f"{'-':>{width}}" if v is None else f"{format(v, spec):>{width}}"


def print_summary(season: str, rows: list[dict], cmp_: Comparison) -> None:
    """Console view of both spans, one line per build."""
    width = max(len(r["label"]) for r in rows) + 2
    print(f"\n[{season}] no supply -> stochastic commitment, {cmp_.what}")
    print(f"  {'build':<{width}}{'p no sup':>10}{'p commit':>10}{'span':>9}"
          f"{'span %':>9}{'profit no sup':>15}{'profit commit':>15}"
          f"{'span':>11}{'span %':>9}")
    for r in rows:
        print(f"  {r['label']:<{width}}"
              f"{_fmt(r['p_none'], '.4f', 10)}{_fmt(r['p_commit'], '.4f', 10)}"
              f"{_fmt(r['price_span'], '+.4f', 9)}"
              f"{_fmt(r['price_span_pct'], '+.1f', 8)}%"
              f"{_fmt(r['profit_none'], ',.2f', 15)}"
              f"{_fmt(r['profit_commit'], ',.2f', 15)}"
              f"{_fmt(r['profit_span'], '+,.2f', 11)}"
              f"{_fmt(r['profit_span_pct'], '+.1f', 8)}%")

    c = contrast(rows, cmp_)
    if c is None:
        return
    print(f"  {'difference':<{width}}"
          f"{_fmt(c['d_p_none'], '+.4f', 10)}{_fmt(c['d_p_commit'], '+.4f', 10)}"
          f"{_fmt(c['d_price_span'], '+.4f', 9)}"
          f"{_fmt(c['d_price_span_pp'], '+.1f', 8)}p"
          f"{_fmt(c['d_profit_none'], '+,.2f', 15)}"
          f"{_fmt(c['d_profit_commit'], '+,.2f', 15)}"
          f"{_fmt(c['d_profit_span'], '+,.2f', 11)}"
          f"{_fmt(c['d_profit_span_pp'], '+.1f', 8)}p")
    print(f"  ({c['against']}; last column in percentage points)")


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def _write_rows(ws, header: list, rows) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    for i, label in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(
            len(str(label)) + 3, 30)


def save_analysis(blobs: dict[str, dict], out_path) -> None:
    """The spans, the build-to-build difference, and every run in full.

    ``blobs`` maps a comparison key to ``{"cmp", "per_season", "runs"}``; both
    comparisons share one workbook, separated by a ``comparison`` column rather
    than by a sheet each, so the four builds can be sorted against each other.
    """
    wb = Workbook()
    wb.remove(wb.active)

    # ---- summary: one block per season, both builds of each comparison ----
    ws = wb.create_sheet("summary")
    header = ["comparison", "season", "build", "DG [kW]", "PV [kWp]",
              "load shift [%]",
              "price, no supply [$/kWh]", "price, stochastic commitment [$/kWh]",
              "price span [$/kWh]", "price span [% of no supply]",
              "price span [x P_ref]",
              "profit, no supply [$/day]", "profit, stochastic commitment [$/day]",
              "profit span [$/day]", "profit span [% of no supply]",
              "E[qMGO], commitment [kWh]", "E[qPG], commitment [kWh]",
              "E[qsupp], commitment [kWh]", "E[qcri], commitment [kWh]"]
    keys = ["label", "dg_kw", "pv_kwp", "shift_pct",
            "p_none", "p_commit", "price_span",
            "price_span_pct", "price_span_x", "profit_none", "profit_commit",
            "profit_span", "profit_span_pct", "qMGO_commit", "qPG_commit",
            "qsupp_commit", "qcri_commit"]
    rows = []
    for key, blob in blobs.items():
        for season, srows in blob["per_season"].items():
            for r in srows:
                rows.append([key, season] + [r[k] for k in keys])
            rows.append([])
    _write_rows(ws, header, rows)

    # ---- what the changed entry did, per comparison and season ----
    ws = wb.create_sheet("effect")
    ckeys = ["comparison", "season", "against", "d_p_none", "d_p_commit",
             "d_price_span", "d_price_span_pp", "d_profit_none",
             "d_profit_commit", "d_profit_span", "d_profit_span_pp"]
    rows = []
    for blob in blobs.values():
        for srows in blob["per_season"].values():
            c = contrast(srows, blob["cmp"])
            if c is not None:
                rows.append([c[k] for k in ckeys])
    _write_rows(
        ws,
        ["comparison", "season", "difference taken",
         "Δ price, no supply [$/kWh]", "Δ price, stochastic commitment [$/kWh]",
         "Δ price span [$/kWh]", "Δ price span [pp]",
         "Δ profit, no supply [$/day]", "Δ profit, stochastic commitment [$/day]",
         "Δ profit span [$/day]", "Δ profit span [pp]"],
        rows)
    ws.append([])
    for blob in blobs.values():
        ws.append([f"{blob['cmp'].key}: {blob['cmp'].fixed}"])
    ws.append(["Spans are already percentages of the no-supply run, so their "
               "differences are in percentage points (pp)."])

    # ---- every run, long form, in the field set the belief sweep uses ----
    ws = wb.create_sheet("runs")
    detail = []
    for key, blob in blobs.items():
        for (season, build_key), b in blob["runs"].items():
            for attr, run_label in RUNS:
                rec = b[attr]
                detail.append([key, season, b["label"], run_label]
                              + [rec[k] for k, _ in FIELDS])
    _write_rows(ws, ["comparison", "season", "build", "run"]
                + [lbl for _, lbl in FIELDS], detail)

    wb.save(out_path)
    print(f"  wrote {out_path}")


# --------------------------------------------------------------------------
# Plot
# --------------------------------------------------------------------------
def _span_rows(per_season: dict[str, list[dict]], lo_key: str, hi_key: str,
               pct_key: str) -> tuple[list[dict], list[str]]:
    """
    One drawable range per (season, build), plus the ones with no value.

    Grouped by season and, within a season, in the comparison's own build
    order, so the two rows of a pair sit next to each other on the figure.
    """
    rows, empty = [], []
    for season, srows in per_season.items():
        for r in srows:
            lo, hi = r[lo_key], r[hi_key]
            if lo is None or hi is None:
                empty.append(f"{season} / {r['label']}")
                continue
            rows.append({"season": season, "build": r["build"],
                         "label": r["label"], "start": lo, "end": hi,
                         "pct": r[pct_key]})
    return rows, empty


def _plot_span(blob: dict, out_path, *, lo_key: str, hi_key: str,
               pct_key: str, xlabel: str, quantity: str, fmt: str,
               marks: list[tuple] = ()) -> None:
    """
    One horizontal range per (season, build): no supply -> stochastic commitment.

    Each line STARTS at the value the leader lands on when EDL is off all day
    (open circle) and RUNS to the value it commits to under the real scenario
    probabilities (filled diamond). Its length is the span, its direction says
    which way EDL's expected presence pushes the quantity, and the two rows of
    a season differ in one config entry only -- so that entry's effect is the
    difference between two line lengths, read directly off the figure.

    Each line is labelled with the span twice: in its own unit, and as a
    percentage of the open-circle end it starts from. The percentage is what
    carries across the price figure and the profit figure, whose units are
    three orders of magnitude apart.

    ``marks`` are vertical reference lines, each
    ``(value, label, linestyle, colour, linewidth)``.
    """
    cmp_ = blob["cmp"]
    rows, empty = _span_rows(blob["per_season"], lo_key, hi_key, pct_key)
    if not rows:
        print(f"  no {hi_key} value in any build — skipping {out_path}")
        return

    fig, ax = plt.subplots(figsize=(10, 0.85 * len(rows) + 3.4))
    ys = list(range(len(rows) - 1, -1, -1))       # first row at the top

    for row, y in zip(rows, ys):
        color = _INK.get(row["build"], "#444444")
        start, end = row["start"], row["end"]

        ax.hlines(y, start, end, color=color, linewidth=2.6, zorder=2)
        ax.plot([start], [y], marker="o", markersize=8, markerfacecolor="white",
                markeredgecolor=color, markeredgewidth=2.0, zorder=3)
        ax.plot([end], [y], marker="D", markersize=8, markerfacecolor=color,
                markeredgecolor="#222222", markeredgewidth=0.7, zorder=3)

        # A span the price menu cannot resolve collapses the two ends onto one
        # point: the markers still stack, but a second value label there would
        # only overprint the first, so it is dropped and the note says why.
        flat = abs(end - start) <= _TOL
        if flat:
            ax.annotate(fmt.format(end), (end, y), xytext=(12, 0),
                        textcoords="offset points", ha="left", va="center",
                        fontsize=8.5, color="#222222")
        else:
            # Value at each end, always on the outer side of its own marker.
            right = end >= start
            for value_, is_end in ((start, False), (end, True)):
                outward = 1 if (is_end == right) else -1
                ax.annotate(fmt.format(value_), (value_, y),
                            xytext=(10 * outward, 0),
                            textcoords="offset points",
                            ha="left" if outward > 0 else "right", va="center",
                            fontsize=8.5, color="#222222")

        if flat:
            note = "span 0 (both runs on the same menu level)"
        else:
            note = f"span {fmt.format(end - start)}"
            if row["pct"] is not None:
                note += f"   ({row['pct']:+.1f}%)"
        ax.annotate(note, (0.5 * (start + end), y), xytext=(0, 10),
                    textcoords="offset points", ha="center", va="bottom",
                    fontsize=7.5, color="#333333")

    # A hairline between seasons, so the pairs read as pairs.
    for i in range(1, len(rows)):
        if rows[i]["season"] != rows[i - 1]["season"]:
            ax.axhline(ys[i] + 0.5, color="#DDDDDD", linewidth=0.8, zorder=0)

    for value_, label, style, color, lw in marks:
        if value_ is None:
            continue
        ax.axvline(value_, color=color, linewidth=lw, linestyle=style, zorder=0)
        if label:
            ax.annotate(label, (value_, 0.0),
                        xycoords=("data", "axes fraction"), xytext=(4, 6),
                        textcoords="offset points", ha="left", va="bottom",
                        fontsize=8, color=color)

    ax.set_yticks(ys)
    ax.set_yticklabels([f"{r['season']} — {r['label']}" for r in rows])
    ax.set_ylim(-0.75, len(rows) - 0.25)
    ax.set_xlabel(xlabel)
    ax.set_ylabel("Season and microgrid build")
    ax.set_title(f"What EDL's expected presence does to the {quantity}, "
                 f"{cmp_.what}\n"
                 "each line runs from EDL off all day to the stochastic "
                 "commitment over all six outcomes;\n"
                 "% is the span against that EDL-off day", fontsize=11.5)
    ax.grid(axis="x", linestyle=":", alpha=0.5)
    ax.margins(x=0.20)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.tick_params(axis="y", length=0)

    handles = [
        plt.Line2D([], [], linestyle="none", marker="o", markersize=8,
                   markerfacecolor="white", markeredgecolor="#444444",
                   markeredgewidth=2.0, label="no supply, prob = 1"),
        plt.Line2D([], [], linestyle="none", marker="D", markersize=8,
                   markerfacecolor="#444444", markeredgecolor="#222222",
                   label="stochastic commitment, real probabilities"),
    ] + [plt.Line2D([], [], color=b.color, linewidth=2.6, label=b.label)
         for b in cmp_.builds]
    ax.legend(handles=handles, fontsize=8, framealpha=0.9, loc="lower right",
              ncol=2)

    caption = [cmp_.fixed]
    if empty:
        caption.append("no value (MGO sells nothing): " + ", ".join(empty))
    fig.tight_layout(rect=(0, 0.022 * len(caption) + 0.012, 1, 1))
    fig.text(0.012, 0.012, "\n".join(caption), fontsize=7.5, color="#777777",
             va="bottom")

    fig.savefig(out_path, dpi=150)
    plt.close(fig)
    print(f"  wrote {out_path}")


def plot_spans(blob: dict, price_path, profit_path) -> None:
    """The price span and the profit span, built the same way from the same
    runs so the two figures can be read against each other."""
    P_ref = config.CONFIG["P_ref"]
    GREY = "#666666"

    _plot_span(
        blob, price_path, lo_key="p_none", hi_key="p_commit",
        pct_key="price_span_pct",
        xlabel="Leader price $\\tilde p$  [\\$ / kWh]",
        quantity="price", fmt="{:.4f}",
        marks=[(P_ref, f"$P_{{ref}}$ = {P_ref:.3f}", ":", GREY, 1.1)])

    _plot_span(
        blob, profit_path, lo_key="profit_none", hi_key="profit_commit",
        pct_key="profit_span_pct",
        xlabel="Leader profit  [\\$ / day]",
        quantity="profit", fmt="{:,.1f}",
        marks=[(0.0, None, "-", GREY, 0.8)])


# --------------------------------------------------------------------------
def main(seasons: list[str] | None = None, levels: int | None = None,
         dg_kw: float = DG_KW, pv_kwp: float = PV_KWP,
         shift_high: float = SHIFT_HIGH,
         comparisons: list[str] | None = None,
         time_limit: float = TIME_LIMIT_S) -> None:
    seasons = seasons or list(config.SEASONS.keys())
    chosen = comparisons or list(COMPARISONS)
    cmps = [fleet_comparison(dg_kw, pv_kwp) if k == "fleet"
            else shift_comparison(dg_kw, pv_kwp, shift_high) for k in chosen]
    out_dir = config.OUTPUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)

    opts = {"MIPGap": config.CONFIG["MIPGap"], "TimeLimit": time_limit}

    # The fleet is this module's own, not the case study's, so say so once
    # rather than let the row labels imply config was changed.
    print(f"microgrid fleet: {_dg_label(dg_kw)}, PV builds + {pv_kwp:.0f} kWp "
          f"(config.CONFIG['CAP'] = {config.CONFIG['CAP']}, left untouched)")
    print(f"solver: MIPGap {opts['MIPGap']:g}, TimeLimit {time_limit:g} s "
          f"per solve (config's {config.CONFIG['TimeLimit']:g} s aborts the "
          f"raised-shift builds)")

    # Builds shared by two comparisons are solved once, so the count is over
    # the DISTINCT models, not over the rows drawn.
    distinct = {b.stamp for c in cmps for b in c.builds}
    print(f"{len(seasons) * len(distinct) * len(RUNS)} solves: "
          f"{len(seasons)} season(s) x {len(distinct)} distinct build(s) x "
          f"{len(RUNS)} runs (no supply, stochastic commitment)"
          + (f"; {len(cmps)} comparison(s) drawn from them"
             if len(cmps) > 1 else ""))

    cache: dict = {}
    blobs: dict[str, dict] = {}
    for cmp_ in cmps:
        print(f"\n=== {cmp_.key}: {cmp_.what} ===")
        per_season: dict[str, list[dict]] = {}
        runs: dict[tuple[str, str], dict] = {}
        for season in seasons:
            rows = []
            for build in cmp_.builds:
                b = solve_build(season, build, opts, levels, cache)
                if b is None:
                    continue
                runs[(season, build.key)] = b
                rows.append(summarize(b))
            if not rows:
                print(f"[{season}] nothing solved.")
                continue
            per_season[season] = rows
            print_summary(season, rows, cmp_)

        # A build that failed to solve leaves a comparison with one row, which
        # on the figure looks like a finished result. Name it instead.
        short = [f"{s} / {b.label}" for s, rows in per_season.items()
                 for b in cmp_.builds
                 if not any(r["build"] == b.key for r in rows)]
        if short:
            print(f"  !! [{cmp_.key}] no span for: {', '.join(short)} — "
                  f"those rows are missing from the figures and the workbook")

        if not per_season:
            print(f"[{cmp_.key}] no season solved — nothing to draw.")
            continue
        blobs[cmp_.key] = {"cmp": cmp_, "per_season": per_season, "runs": runs}
        plot_spans(blobs[cmp_.key],
                   out_dir / f"stochastic_commitment_price_span_{cmp_.key}.png",
                   out_dir / f"stochastic_commitment_profit_span_{cmp_.key}.png")

    if not blobs:
        print("Nothing solved — no workbook written.")
        return
    save_analysis(blobs, out_dir / "stochastic_commitment_analysis.xlsx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Compare the no-supply-to-commitment span of the leader's "
                    "price and profit between microgrid builds: genset-only "
                    "against genset+PV, and the standard load shift against a "
                    "raised one.")
    parser.add_argument("--seasons", nargs="+",
                        default=list(config.SEASONS.keys()),
                        choices=list(config.SEASONS.keys()),
                        help="Seasons to solve (default: all three).")
    parser.add_argument("--comparisons", nargs="+", default=None,
                        choices=list(COMPARISONS),
                        help="Which comparisons to draw (default: both). They "
                             "share the genset+PV build, which is solved once.")
    parser.add_argument("--dg", type=float, default=DG_KW, metavar="KW",
                        help=f"Diesel capacity of every build in kW (default "
                             f"{DG_KW:g}). Patched for the solves only; "
                             f"config.CONFIG['CAP']['DG'] is left alone.")
    parser.add_argument("--pv", type=float, default=PV_KWP, metavar="KWP",
                        help=f"PV capacity of the PV builds in kWp (default "
                             f"{PV_KWP:g}). The genset-only build has none.")
    parser.add_argument("--shift", type=float, default=SHIFT_HIGH,
                        metavar="FRAC",
                        help=f"Raised load shift of the shift comparison, as a "
                             f"fraction of each hour's demand (default "
                             f"{SHIFT_HIGH:g}). The other build keeps the "
                             f"configured "
                             f"{config.CONFIG['shift_frac']:g}.")
    parser.add_argument("--time-limit", type=float, default=TIME_LIMIT_S,
                        metavar="S",
                        help=f"Solver time limit per solve in seconds "
                             f"(default {TIME_LIMIT_S:g}). config's "
                             f"{config.CONFIG['TimeLimit']:g} s is too short "
                             f"for the raised-shift builds, which abort and "
                             f"drop their row.")
    parser.add_argument("--levels", type=int, default=None,
                        help="Re-space the price menu to this many levels; "
                             "default keeps the configured menu. The menu step "
                             "is the resolution of the price span this module "
                             "reports.")
    args = parser.parse_args()
    if args.dg <= 0.0:
        parser.error("--dg must be positive: the microgrid needs a genset.")
    if args.pv <= 0.0:
        parser.error("--pv must be positive: with no PV the two builds of the "
                     "fleet comparison are the same one.")
    if not 0.0 < args.shift <= 1.0:
        parser.error("--shift must lie in (0, 1]: it is a fraction of each "
                     "hour's demand.")
    main(args.seasons, args.levels, args.dg, args.pv, args.shift,
         args.comparisons, args.time_limit)
