"""
How the leader's price answers HOW MANY hours EDL supplies, and which ones.

:mod:`pattern_sweep` fixes EDL's window at 8 hours and asks whether it sits on
the highest- or the lowest-demand hours. This module keeps that question and
sweeps the window's length instead: EDL supplies the N highest-demand hours
(peak) or the N lowest (offpeak) of each season, for every N from 1 to 23,
and nothing in the rest of the day.

    peak      the N hours of highest village net demand
    offpeak   the N hours of lowest village net demand

Every run is deterministic -- P(pattern) = 1 -- so the leader knows the window
when it posts its price.

THE HOURS, TAKEN OVER COMPLETELY                                 [ASSUMPTION]
As in :mod:`pattern_sweep`, EDL's cap in a supplied hour is ``M_phys``: it covers
whatever demand is there, including whatever the village shifts in. Longer
windows therefore deliver more energy, and peak windows more than offpeak ones
of the same length; the workbook carries each window's share of daily demand.

WHAT IS REPORTED: THE PRICE DEVIATION FROM NO SUPPLY
    deviation   = p(window) - p(no supply)              [$/kWh]
    deviation % = 100 * deviation / p(no supply)        [% of no-supply price]

The reference is the same season, fleet and load shift with EDL off all day, so
N = 0 IS the reference and every line starts at zero there.

THE REGRET OF KEEPING THE NO-SUPPLY PRICE                          (--regret)
    regret   = profit - profit at p(no supply)            [$/day]
    regret % = 100 * regret / profit                      [% of optimal profit]

"Profit at p(no supply)" is the same window, the village scheduling for it,
with the price menu reduced to that season's no-supply price: what the leader
earns if it ignores EDL's window and keeps its no-supply price. That price is
on the full menu, so the regret is non-negative up to the solver gap (any
value below it is printed); it is 0 at 0 h and wherever the two prices
coincide, which needs no solve.

THE FLEET                                                        [ASSUMPTION]
Fixed: config's 400 kW genset plus 200 kWp of PV, at 10% load shift. The PV is
what gives the hours different costs, so the window's placement can move the
price through the cost channel described in :mod:`pattern_sweep`, not only
through load shifting. Only the window is swept.

``--no-pv`` runs the same sweep on the genset alone. Every served hour then
costs the same, the cost channel is gone, and whatever the window still does to
the price comes through the load the village shifts into the covered hours.

Outputs, under :data:`config.OUTPUT_DIR`:

    hour_sweep.png     2 x 2: peak on top, offpeak below; the price left, its
                       % deviation from no supply right; one line per season
    hour_sweep.xlsx    every run behind it
    hour_sweep_nopv.*  the same two, for --no-pv
    hour_sweep_profit.png  2 x 2: the leader's profit, no PV left and PV
                       right, peak on top and offpeak below (--profit)
    hour_sweep_regret.png  the same panels: the regret of keeping the
                       no-supply price, in % of optimal profit (--regret;
                       the regret columns are added to both workbooks)

Usage:
    python hour_sweep.py
    python hour_sweep.py --no-pv
    python hour_sweep.py --profit          # needs both workbooks; no solving
    python hour_sweep.py --regret          # needs both workbooks; fixed-price solves
    python hour_sweep.py --regret --replot # redraw the regret figure only
    python hour_sweep.py --seasons winter
    python hour_sweep.py --replot          # redraw, no solving
"""

import argparse
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from pyomo.environ import TerminationCondition, value

import config
from asset_sweep import (INK, SEASON_COLOR, SEASON_LABEL, SEASON_MARKER,
                         AXIS, MUTED, RULE, _layout_issues, _style,
                         _write_rows,
                         menu_step)
from model.bilevel_mpec import BilevelMPEC
from pattern_sweep import solve_at_belief
from run import build_instance
from stochastic_commitment_analysis import TIME_LIMIT_S, Build, build_settings

# Window lengths swept [h]: every hour count.                     [ASSUMPTION]
# 23 h leaves the microgrid only the one hour EDL does not cover; 24 h would
# leave it nothing to sell and no price, so the sweep stops there.
HOURS: tuple[int, ...] = tuple(range(1, 24))

PATTERNS: tuple[str, ...] = ("peak", "offpeak")
NONE = "no supply"

# The fixed build: config's genset, 200 kWp of PV, 10% load shift. [ASSUMPTION]
DG_KW: float = config.CONFIG["CAP"]["DG"]
PV_KWP: float = 200.0
SHIFT: float = 0.10

FIELDS = [
    ("season",       "season"),
    ("fleet",        "fleet"),
    ("shift",        "load shift [-]"),
    ("series",       "series"),
    ("hours",        "hours supplied [h]"),
    ("on_hours",     "supplied hours (hour-beginning)"),
    ("share",        "demand in the supplied hours / daily demand [-]"),
    ("p",            "leader price [$/kWh]"),
    ("p_ref",        "price, no supply [$/kWh]"),
    ("d_p",          "price deviation [$/kWh]"),
    ("d_p_pct",      "price deviation [% of no supply]"),
    ("d_p_steps",    "price deviation [menu steps]"),
    ("profit",       "leader profit [$/day]"),
    ("profit_ref",   "profit, no supply [$/day]"),
    ("d_profit",     "profit deviation [$/day]"),
    ("profit_hold",  "profit at the no-supply price [$/day]"),
    ("regret",       "regret of keeping the no-supply price [$/day]"),
    ("regret_pct",   "regret of keeping the no-supply price [% of profit]"),
    ("qMGO",        "microgrid sales [kWh]"),
    ("qPG",          "EDL delivered [kWh]"),
    ("qsupp",        "suppressed demand [kWh]"),
    ("qcri",         "unserved [kWh]"),
    ("termination",  "termination"),
]


# --------------------------------------------------------------------------
# Solving
# --------------------------------------------------------------------------
def supplied_hours(D: dict[int, float], n: int, pattern: str) -> list[int]:
    """The ``n`` highest- (peak) or lowest-demand (offpeak) hours, in clock
    order."""
    ranked = sorted(D, key=lambda t: D[t], reverse=(pattern == "peak"))
    return sorted(ranked[:n])


def _fleet_label(dg_kw: float, pv_kwp: float) -> str:
    return (f"{dg_kw:.0f} kW DG" if pv_kwp <= 0.0
            else f"{dg_kw:.0f} kW DG + {pv_kwp:.0f} kWp PV")


def _record(rec: dict, ref: dict, step: float, **extra) -> dict:
    """A solved run with its deviation from the no-supply reference."""
    out = dict(rec) | extra
    out["p_ref"], out["profit_ref"] = ref["p"], ref["profit"]
    p, p0 = out["p"], ref["p"]
    out["d_p"] = None if p is None or p0 is None else p - p0
    out["d_p_pct"] = None if out["d_p"] is None else 100.0 * out["d_p"] / p0
    out["d_p_steps"] = None if out["d_p"] is None else out["d_p"] / step
    out["d_profit"] = out["profit"] - ref["profit"]
    return out


def sweep(seasons: list[str], hours: list[int], build: Build, opts: dict,
          levels: int | None, step: float) -> list[dict]:
    """The no-supply reference and both patterns at every window length."""
    cap = config.CONFIG["M_phys"]
    rows = []
    for season in seasons:
        D = build_instance(season)["D"]
        T = sorted(D)
        total = sum(D.values())
        print(f"\n[{season}] {build.label}")

        with build_settings(build):
            ref = solve_at_belief(season, {t: 0.0 for t in T}, 1.0, opts,
                                  levels)
        if ref is None or ref["p"] is None:
            print(f"[{season}] no-supply reference did not solve — skipped")
            continue
        common = {"season": season, "fleet": build.label,
                  "shift": build.shift_frac}
        rows.append(_record(ref, ref, step, **common, series=NONE, hours=0,
                            on_hours="", share=0.0))
        print(f"    {NONE:>10}  0 h  p={ref['p']:.4f}  "
              f"profit={ref['profit']:8.2f}")

        for pattern in PATTERNS:
            for n in hours:
                on = supplied_hours(D, n, pattern)
                Qbar = {t: (cap if t in on else 0.0) for t in T}
                with build_settings(build):
                    rec = solve_at_belief(season, Qbar, 1.0, opts, levels)
                if rec is None:
                    print(f"    {pattern:>10} {n:2d} h  did not solve — "
                          f"missing from the figure")
                    continue
                row = _record(rec, ref, step, **common, series=pattern,
                              hours=n, on_hours=",".join(map(str, on)),
                              share=sum(D[t] for t in on) / total)
                rows.append(row)
                dp = ("   -   " if row["d_p"] is None
                      else f"{row['d_p']:+.4f} ({row['d_p_pct']:+.1f}%)")
                print(f"    {pattern:>10} {n:2d} h  hours {row['on_hours']:<30}"
                      f"  Δp={dp}  qPG={row['qPG']:7.1f}")
    return rows


# --------------------------------------------------------------------------
# Solving: the regret of keeping the no-supply price
# --------------------------------------------------------------------------
_OK = (TerminationCondition.optimal, TerminationCondition.feasible)


def solve_at_price(season: str, Qbar: dict, price: float,
                   opts: dict) -> dict | None:
    """
    The certain window with the menu reduced to ``price``: the leader cannot
    re-price, the village schedules for the window it knows.
    """
    data = dict(build_instance(season))
    T = data["T"]
    data["S"], data["prob"] = ["window"], {"window": 1.0}
    data["QbarPG"] = {(t, "window"): Qbar[t] for t in T}
    data["Pi"] = [price]

    mpec = BilevelMPEC(data)
    m = mpec.build()
    solver_name, res = mpec.solve("gurobi", options=opts)
    if solver_name is None:
        raise RuntimeError(
            "Gurobi not available (check the license and that gurobipy is "
            "installed).")
    tc = res.solver.termination_condition
    if tc not in _OK:
        print(f"    not solved ({tc})")
        return None
    return {"profit": value(m.profit), "termination": str(tc)}


def regret_runs(rows: list[dict], build: Build, opts: dict) -> None:
    """
    Fill in the profit at the no-supply price and the regret of every run that
    lacks them, in place. Runs at the no-supply price already -- 0 h, and
    every window whose optimum is that price -- take their own profit.
    """
    cap = config.CONFIG["M_phys"]
    tol = config.CONFIG["MIPGap"]
    for season in [s for s in config.SEASONS
                   if any(r["season"] == s for r in rows)]:
        T = sorted(build_instance(season)["D"])
        todo = [r for r in rows if r["season"] == season
                and r.get("regret_pct") is None and r.get("d_p") is not None]
        if not todo:
            continue
        print(f"\n[{season}] {build.label}, regret of keeping the no-supply "
              f"price ({len(todo)} windows)")
        for r in sorted(todo, key=lambda r: (r["series"], r["hours"])):
            if abs(r["d_p"]) <= 1e-9:
                hold = r["profit"]
            else:
                on = {int(h) for h in str(r.get("on_hours") or "").split(",")
                      if h.strip()}
                Qbar = {t: (cap if t in on else 0.0) for t in T}
                with build_settings(build):
                    rec = solve_at_price(season, Qbar, r["p_ref"], opts)
                if rec is None:
                    continue
                hold = rec["profit"]
            r["profit_hold"] = hold
            r["regret"] = r["profit"] - hold
            r["regret_pct"] = (100.0 * r["regret"] / r["profit"]
                               if r["profit"] > 1e-9 else None)
            print(f"    {r['series']:>10} {r['hours']:2d} h  "
                  f"Δp={r['d_p_steps']:+3.0f} steps  profit={r['profit']:8.2f}  "
                  f"at p(no supply)={hold:8.2f}  regret={r['regret']:6.2f} "
                  f"({r['regret_pct'] or 0.0:5.2f}%)")
            if r["regret"] < -max(1e-6, tol * abs(r["profit"])):
                print(f"    !! negative regret: {season}, {r['series']}, "
                      f"{r['hours']} h — the optimal run fell short")


# --------------------------------------------------------------------------
# Plot
# --------------------------------------------------------------------------
COLUMNS = (
    ("p", "Leader price   [$ / kWh]"),
    ("d_p_pct", "Price deviation from no supply   [% of no-supply price]"),
)
ROW_TITLE = {"peak": "EDL on the highest-demand hours (peak)",
             "offpeak": "EDL on the lowest-demand hours (offpeak)"}


def _chrome(ax, zero: bool) -> None:
    """
    :func:`asset_sweep._style`, whose zero line is kept only where 0 is a
    value on the axis -- on a price level it would stretch the axis down to 0.
    """
    if zero:
        _style(ax)
        return
    ax.grid(axis="y", color=RULE, linewidth=0.7)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(AXIS)
    ax.tick_params(colors=MUTED, labelsize=9)


def _season_lines(ax, rows: list[dict], pattern: str, key: str,
                  order: list[str]) -> None:
    """
    One line per season of ``key`` against the window length. N = 0 is the
    no-supply reference itself, so every line starts there.
    """
    for season in order:
        pts = sorted((r["hours"], r[key]) for r in rows
                     if r["season"] == season and r.get(key) is not None
                     and (r["series"] == pattern
                          or (r["series"] == NONE and r["hours"] == 0)))
        if len(pts) < 2:
            continue
        xs, ys = zip(*pts)
        # Smaller markers once every hour count is solved, so 24 points on
        # one line still read as a line.
        ax.plot(xs, ys, color=SEASON_COLOR[season], linewidth=2.0,
                marker=SEASON_MARKER[season],
                markersize=6.5 if len(xs) <= 12 else 4.5,
                markeredgecolor="white", markeredgewidth=1.0, zorder=3)


def _finish(fig, axes, out_path, *, order: list[str], hours: list[int],
            title: str) -> None:
    """Shared x ticks and x labels, the season legend, the title, and save."""
    for ax in axes[-1]:
        ax.set_xlabel("Hours EDL supplies   [h]", fontsize=10, color=INK)
    # Shared x hides the upper row's hour labels; every panel keeps them.
    for ax in axes.flat:
        ax.tick_params(axis="x", labelbottom=True)
    swept = sorted({0, *hours})
    # Every solved length gets a tick while they fit; an hour-by-hour sweep is
    # labelled every second hour so the axis stays readable.
    ticks = swept if len(swept) <= 13 else list(range(0, max(swept) + 1, 2))
    axes[0, 0].set_xticks(ticks)
    axes[0, 0].set_xlim(-0.5, max(swept) + 0.5)

    handles = [plt.Line2D([], [], color=SEASON_COLOR[s], linewidth=2.0,
                          marker=SEASON_MARKER[s], markersize=6.5,
                          markeredgecolor="white",
                          label=SEASON_LABEL.get(s, s)) for s in order]
    fig_w = fig.get_size_inches()[0]
    ttl = fig.suptitle(textwrap.fill(title, width=int(8.2 * fig_w)),
                       x=0.01, y=0.985, ha="left", va="top", fontsize=12,
                       color=INK, fontweight="bold")
    ttl_lines = ttl.get_text().count("\n")
    leg = fig.legend(handles=handles, loc="upper center",
                     bbox_to_anchor=(0.5, 0.965 - 0.032 * (ttl_lines + 1)),
                     ncol=len(handles), frameon=False, fontsize=9)
    fig.tight_layout(rect=(0, 0.01, 1, 0.935 - 0.032 * (ttl_lines + 1)))
    for msg in _layout_issues(fig, [ttl, leg]):
        print(f"  layout: {msg}  [{out_path}]")
    fig.savefig(out_path, dpi=150, facecolor="white")
    plt.close(fig)
    print(f"  wrote {out_path}")


def plot(rows: list[dict], out_path, *, seasons: list[str], hours: list[int],
         step: float, fleet: str, shift: float) -> None:
    """
    Four panels: peak on top, offpeak below; the leader's price on the left and
    its deviation in % of the no-supply price on the right. Each column shares
    its y axis, so the two patterns are read against each other on one scale.
    """
    order = [s for s in config.SEASONS if s in seasons]
    fig, axes = plt.subplots(2, 2, figsize=(13.0, 8.8), sharex=True,
                             sharey="col")

    for i, pattern in enumerate(PATTERNS):
        for j, (key, label) in enumerate(COLUMNS):
            ax = axes[i, j]
            _chrome(ax, zero=(key != "p"))
            _season_lines(ax, rows, pattern, key, order)
            if key == "p":
                # The elasticity anchor, which is also the bottom of the price
                # menu, so the levels read against it.
                p_ref = config.CONFIG["P_ref"]
                ax.axhline(p_ref, color=MUTED, linewidth=1.0,
                           linestyle=(0, (4, 3)), zorder=2)
                ax.annotate(f"$P_{{ref}}$ = {p_ref:.4f}", (0.0, p_ref),
                            xycoords=("axes fraction", "data"),
                            xytext=(4, 3), textcoords="offset points",
                            ha="left", va="bottom", fontsize=8.5,
                            color=MUTED)
            ax.set_title(ROW_TITLE[pattern], loc="left", fontsize=10.5,
                         color=INK)
            ax.set_ylabel(label, fontsize=9.5, color=INK)

    axes[0, 0].yaxis.set_major_formatter(lambda v, _: f"{v:.3f}")
    axes[0, 1].yaxis.set_major_formatter(lambda v, _: f"{v:+.1f}%"
                                         if abs(v) > 1e-12 else "0%")
    _finish(fig, axes, out_path, order=order, hours=hours,
            title=f"The leader's price against the length of EDL's window, "
                  f"on {fleet} at {100 * shift:.0f}% load shift. EDL covers "
                  f"its hours completely and nothing else; 0 h is the "
                  f"no-supply run each line is measured from")


def plot_profit(by_fleet: list[tuple[str, list[dict]]], out_path, *,
                seasons: list[str], shift: float) -> None:
    """
    The leader's profit, one column per fleet: peak on top, offpeak below. All
    four panels share one y axis, so the gap between the columns is what the
    PV is worth to the leader at each window length.
    """
    _plot_by_fleet(
        by_fleet, out_path, seasons=seasons, key="profit",
        ylabel="Leader profit   [$ / day]", fmt=lambda v, _: f"{v:,.0f}",
        title=f"The leader's profit against the length of EDL's window, "
              f"without PV (left) and with PV (right), at "
              f"{100 * shift:.0f}% load shift. EDL covers its hours "
              f"completely and nothing else; 0 h is the no-supply run")


def plot_regret(by_fleet: list[tuple[str, list[dict]]], out_path, *,
                seasons: list[str], shift: float) -> None:
    """The regret of keeping the no-supply price, in % of optimal profit, in
    the panels of :func:`plot_profit`; each line's peak carries its absolute
    regret in $/day."""
    _plot_by_fleet(
        by_fleet, out_path, seasons=seasons, key="regret_pct",
        ylabel="Regret   [% of optimal profit]",
        fmt=lambda v, _: f"{v:g}%",
        peak_label=("regret", lambda v: f"{v:.1f} $/day"),
        title=f"What keeping the no-supply price would cost the leader, "
              f"against the length of EDL's window, without PV (left) and "
              f"with PV (right), at {100 * shift:.0f}% load shift. Regret = "
              f"optimal profit minus profit at the no-supply price, the "
              f"village scheduling for the window in both")


def _label_peaks(ax, rows: list[dict], pattern: str, key: str,
                 order: list[str], label_key: str, label_fmt) -> None:
    """
    Ring each season line's highest point and write ``label_key`` of that run
    beside it, in the season's colour. Labels are stacked by season order so
    that peaks at the same height do not print over each other.
    """
    for k, season in enumerate(order):
        sel = [r for r in rows if r["season"] == season
               and r["series"] == pattern and r.get(key) is not None
               and r.get(label_key) is not None]
        if not sel:
            continue
        # A flat top (the same % at several lengths) is labelled where the
        # labelled field is largest, not at whichever length comes first.
        top = max(sel, key=lambda r: (round(r[key], 6), r[label_key]))
        color = SEASON_COLOR[season]
        ax.plot(top["hours"], top[key], marker="o", markersize=11,
                markerfacecolor="none", markeredgecolor=color,
                markeredgewidth=1.4, zorder=4)
        ax.annotate(label_fmt(top[label_key]), (top["hours"], top[key]),
                    xytext=(0, 9 + 12 * k), textcoords="offset points",
                    ha="center", va="bottom", fontsize=8.5, color=color,
                    fontweight="bold", zorder=5,
                    bbox=dict(boxstyle="round,pad=0.15", facecolor="white",
                              edgecolor="none", alpha=0.85))


def _plot_by_fleet(by_fleet: list[tuple[str, list[dict]]], out_path, *,
                   seasons: list[str], key: str, ylabel: str, fmt,
                   title: str, peak_label=None) -> None:
    """``key`` against the window length, one column per fleet: peak on top,
    offpeak below, all four panels on one y axis. ``peak_label`` =
    ``(field, format)`` labels each line's peak with that field."""
    order = [s for s in config.SEASONS if s in seasons]
    hours = sorted({r["hours"] for _, rows in by_fleet for r in rows
                    if r["hours"]})
    fig, axes = plt.subplots(2, len(by_fleet), figsize=(13.0, 8.8),
                             sharex=True, sharey=True, squeeze=False)

    for i, pattern in enumerate(PATTERNS):
        for j, (fleet, rows) in enumerate(by_fleet):
            ax = axes[i, j]
            _style(ax)
            # Shared axes hide their inner tick labels; each column is read on
            # its own, so every panel keeps its numbers.
            ax.tick_params(axis="y", labelleft=True)
            _season_lines(ax, rows, pattern, key, order)
            if peak_label:
                _label_peaks(ax, rows, pattern, key, order, *peak_label)
            ax.set_title(f"{ROW_TITLE[pattern]}  ·  {fleet}", loc="left",
                         fontsize=10, color=INK)
            if j == 0:
                ax.set_ylabel(ylabel, fontsize=9.5, color=INK)

    if peak_label:
        # Headroom for the stacked labels above the highest peak.
        lo, hi = axes[0, 0].get_ylim()
        axes[0, 0].set_ylim(lo, hi + 0.22 * (hi - lo))
    axes[0, 0].yaxis.set_major_formatter(fmt)
    _finish(fig, axes, out_path, order=order, hours=hours, title=title)


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def save_analysis(rows: list[dict], out_path, step: float) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "runs"
    _write_rows(ws, [lbl for _, lbl in FIELDS],
                ([r.get(k) for k, _ in FIELDS] for r in rows))
    ws.append([])
    ws.append([f"Deterministic runs, P(window) = 1. EDL's cap is M_phys in the "
               f"supplied hours and 0 elsewhere. Deviations are against the "
               f"no-supply row of the same season. Price-menu step "
               f"{step:.5f} $/kWh."])
    wb.save(out_path)
    print(f"  wrote {out_path}")


def load_rows(path) -> list[dict]:
    """Every run of a previous sweep, read back from its workbook."""
    ws = load_workbook(path, data_only=True)["runs"]
    table = list(ws.iter_rows(values_only=True))
    key_of = {label: key for key, label in FIELDS}
    header = [key_of.get(cell) for cell in table[0]]
    recs = [{k: v for k, v in zip(header, row) if k is not None}
            for row in table[1:]]
    # The blank line and the caption carry no series.
    return [r for r in recs if r.get("series") is not None]


# --------------------------------------------------------------------------
def main(seasons: list[str] | None = None, time_limit: float = TIME_LIMIT_S,
         replot: bool = False, no_pv: bool = False,
         profit: bool = False, regret: bool = False) -> None:
    seasons = seasons or list(config.SEASONS.keys())
    hours = list(HOURS)
    shift = SHIFT
    pv_kwp = 0.0 if no_pv else PV_KWP
    fleet = _fleet_label(DG_KW, pv_kwp)
    levels = None
    step = menu_step(levels)
    out_dir = config.OUTPUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = "hour_sweep_nopv" if no_pv else "hour_sweep"
    book, figure = out_dir / f"{stem}.xlsx", out_dir / f"{stem}.png"

    if regret:
        # Both fleets side by side. The optimal runs come from the workbooks;
        # only the fixed-price runs they lack are solved, then written back.
        books = [(out_dir / "hour_sweep_nopv.xlsx", 0.0),
                 (out_dir / "hour_sweep.xlsx", PV_KWP)]
        missing = [str(b) for b, _ in books if not b.exists()]
        if missing:
            print(f"missing {', '.join(missing)} — run the sweep for that "
                  f"fleet first")
            return
        opts = {"MIPGap": config.CONFIG["MIPGap"], "TimeLimit": time_limit}
        by_fleet = []
        for b, pv in books:
            rows = load_rows(b)
            for r in rows:
                if r["series"] == NONE and r.get("regret_pct") is None:
                    r["profit_hold"], r["regret"], r["regret_pct"] = (
                        r["profit"], 0.0, 0.0)
            if not replot:
                build = Build("regret", rows[0]["fleet"], DG_KW, pv,
                              rows[0]["shift"])
                regret_runs([r for r in rows if r["season"] in seasons],
                            build, opts)
                save_analysis(rows, b, step)
            by_fleet.append((rows[0]["fleet"], rows))
            print(f"\n=== largest regret of keeping the no-supply price, "
                  f"{rows[0]['fleet']}")
            for season in [s for s in config.SEASONS if s in seasons]:
                for pattern in PATTERNS:
                    sel = [r for r in rows if r["season"] == season
                           and r["series"] == pattern
                           and r.get("regret_pct") is not None]
                    if sel:
                        top = max(sel, key=lambda r: r["regret_pct"])
                        print(f"  {season:>12} {pattern:>8}  "
                              f"{top['regret_pct']:5.2f}% "
                              f"({top['regret']:6.2f} $/day) at "
                              f"{top['hours']:2d} h, Δp "
                              f"{top['d_p_pct']:+5.1f}%")
        plot_regret(by_fleet, out_dir / "hour_sweep_regret.png",
                    seasons=seasons, shift=by_fleet[0][1][0]["shift"])
        return

    if profit:
        # Both fleets side by side, from their workbooks; no solving.
        books = [out_dir / "hour_sweep_nopv.xlsx", out_dir / "hour_sweep.xlsx"]
        missing = [str(b) for b in books if not b.exists()]
        if missing:
            print(f"missing {', '.join(missing)} — run the sweep for that "
                  f"fleet first")
            return
        by_fleet = []
        for b in books:
            rows = load_rows(b)
            print(f"read {b}: {len(rows)} runs")
            by_fleet.append((rows[0]["fleet"], rows))
        plot_profit(by_fleet, out_dir / "hour_sweep_profit.png",
                    seasons=seasons, shift=by_fleet[0][1][0]["shift"])
        return

    if replot:
        rows = load_rows(book)
        print(f"replot from {book}: {len(rows)} runs")
        if rows:
            fleet, shift = rows[0]["fleet"], rows[0]["shift"]
        hours = sorted({r["hours"] for r in rows if r["hours"]})
        plot(rows, figure, seasons=seasons, hours=hours, step=step,
             fleet=fleet, shift=shift)
        return

    print(f"{len(seasons) * (1 + len(PATTERNS) * len(hours))} solves: "
          f"{len(seasons)} season(s), windows of {', '.join(map(str, hours))} h")
    opts = {"MIPGap": config.CONFIG["MIPGap"], "TimeLimit": time_limit}
    build = Build("hours", fleet, DG_KW, pv_kwp, shift)
    rows = sweep(seasons, hours, build, opts, levels, step)
    if not rows:
        print("Nothing solved — no figure or workbook written.")
        return
    save_analysis(rows, book, step)
    plot(rows, figure, seasons=seasons, hours=hours, step=step, fleet=fleet,
         shift=shift)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Price deviation from no supply when EDL covers the N "
                    "highest- or lowest-demand hours, across N.")
    parser.add_argument("--seasons", nargs="+",
                        default=list(config.SEASONS.keys()),
                        choices=list(config.SEASONS.keys()))
    parser.add_argument("--time-limit", type=float, default=TIME_LIMIT_S,
                        metavar="S", help="Solver time limit per solve.")
    parser.add_argument("--replot", action="store_true",
                        help="Redraw the figure from the existing workbook "
                             "without solving anything.")
    parser.add_argument("--no-pv", action="store_true",
                        help=f"Genset only, instead of {PV_KWP:g} kWp of PV; "
                             f"writes hour_sweep_nopv.*.")
    parser.add_argument("--profit", action="store_true",
                        help="Draw the profit figure, no PV against PV, from "
                             "both workbooks without solving anything.")
    parser.add_argument("--regret", action="store_true",
                        help="Solve every window at its season's no-supply "
                             "price and draw the regret of keeping that "
                             "price, both fleets, from both workbooks; with "
                             "--replot only redraw.")
    args = parser.parse_args()
    main(args.seasons, args.time_limit, args.replot, args.no_pv, args.profit,
         args.regret)
