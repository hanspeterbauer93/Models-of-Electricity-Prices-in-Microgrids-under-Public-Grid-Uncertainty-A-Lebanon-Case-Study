"""
Scenario mismatch analysis: price for one EDL outcome, live with another.

The leader commits ONE price schedule before the public-grid (EDL) scenario
realizes -- in :mod:`model.bilevel_mpec` only ``p`` and the menu binaries
``b`` are first stage, everything else is recourse. This module asks what that
commitment costs when the belief is wrong:

    row i  ("commitment")   the leader solves the model as if outcome i were
                            CERTAIN (prob = 1) and commits the resulting price
                            schedule p^i_t
    col j  ("realization")  outcome j actually happens: p is held FIXED at p^i,
                            the follower re-optimizes through its full KKT
                            system and the MGO redispatches under j

    profit[i][j]            the leader profit that combination realizes
    regret[i][j]            profit[j][j] - profit[i][j] >= 0, i.e. what the
                            wrong belief cost relative to perfect foresight

The diagonal is perfect foresight (belief = realization). One extra row,
``stochastic``, is the stochastic commitment from the full six-scenario solve -- the
schedule the model actually recommends, evaluated outcome by outcome.

Weighting each row by the scenario probabilities gives:

    E[profit | commitment i]     expected profit of committing to p^i
    E[profit | perfect foresight] = sum_j prob_j * profit[j][j]
    EVPI = E[perfect foresight] - E[stochastic]      value of knowing EDL
    cost of belief i = E[stochastic] - E[i]          price of guessing instead
                                                     of committing
                                                     stochastically

Outputs, under :data:`config.OUTPUT_DIR`:

    scenario_analysis.xlsx              profit / regret / price matrices, the
                                        full per-cell detail and a summary
    scenario_mismatch_<season>.png      profit and regret heatmaps
    scenario_expected_<season>.png      expected profit per commitment, against
                                        the stochastic-commitment and
                                        perfect-foresight bars

Usage:
    python scenario_analysis.py
    python scenario_analysis.py --seasons winter --no-stochastic
"""

import argparse

import matplotlib
matplotlib.use("Agg")             # headless / file output
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap, Normalize, TwoSlopeNorm
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pyomo.environ import value, TerminationCondition

import config
from model.bilevel_mpec import BilevelMPEC
# The stochastic-commitment key/label and the label lookup are shared with :mod:`plot`, so the
# extra row here and the extra box there are named identically.
from plot import COLORS, STOCHASTIC, label_of
from run import build_instance
from scenarios import SCENARIOS, SCENARIO_PROB

_OK = (TerminationCondition.optimal, TerminationCondition.feasible)

# Per-cell fields, in the order they are written to the detail sheet.
FIELDS = [
    ("commitment",    "commitment (priced for)"),
    ("realization",   "realization (happened)"),
    ("profit",        "leader profit [$]"),
    ("regret",        "regret vs perfect foresight [$]"),
    ("revenue",       "leader revenue [$]"),
    ("gen_cost",      "generation cost [$]"),
    ("follower_cost", "follower cost [$]"),
    ("qMGO",          "sum qMGO [kWh]"),
    ("qPG",           "sum qPG [kWh]"),
    ("qcri",          "sum qcri [kWh]"),
    ("qsupp",         "sum qsupp [kWh]"),
    ("D",             "sum D [kWh]"),
    ("supp_share",    "qsupp / D [%]"),
    ("p_mean",        "mean price [$/kWh]"),
    ("termination",   "termination"),
]


# --------------------------------------------------------------------------
# Solving
# --------------------------------------------------------------------------
def _instance(season: str, scenario: str | None) -> dict:
    """Instance data for ``season``; with ``scenario`` given, that EDL outcome
    is treated as CERTAIN (probability 1, the others dropped)."""
    data = build_instance(season)
    if scenario is None:
        return data
    T = data["T"]
    data["S"] = [scenario]
    data["prob"] = {scenario: 1.0}
    data["QbarPG"] = {(t, scenario): data["QbarPG"][(t, scenario)] for t in T}
    return data


def _pin_price(data: dict, p_sched: dict[int, float]) -> dict[int, int]:
    """
    Collapse the price menu onto a committed schedule.

    The menu ``Pi`` is shared by all hours, so a single-entry menu (as the
    archived price sweep used) can only pin a FLAT price. A committed schedule
    varies hour to hour, so the menu is set to its distinct levels and the
    caller fixes ``b[t, k]`` to the level this hour committed to.

    Returns ``{hour: level index}``.
    """
    T = data["T"]
    levels = sorted({round(p_sched[t], 12) for t in T})
    data["Pi"] = levels
    # The cap only bounds p[t]; keeping it at the highest committed level lets
    # every hour hold its own price.
    data["Phat"] = {t: levels[-1] for t in T}
    return {t: levels.index(round(p_sched[t], 12)) for t in T}


def _solve(mpec: BilevelMPEC, opts: dict):
    """Solve and return the termination condition (raises if Gurobi is absent)."""
    solver_name, res = mpec.solve("gurobi", options=opts)
    if solver_name is None:
        raise RuntimeError(
            "Gurobi not available (check the license and that gurobipy is "
            "installed)."
        )
    return res.solver.termination_condition


def _metrics(m, data: dict) -> dict:
    """Leader / follower totals of a solved single-outcome model."""
    T, S, prob = data["T"], data["S"], data["prob"]

    def exp_sum(var):
        return sum(prob[s] * value(var[t, s]) for t in T for s in S)

    qMGO, qPG, qcri = exp_sum(m.qMGO), exp_sum(m.qPG), exp_sum(m.qcri)
    qsupp = exp_sum(m.qsupp)
    sD = sum(data["D"][t] for t in T)

    revenue = sum(prob[s] * value(m.p[t]) * value(m.qMGO[t, s])
                  for t in T for s in S)
    gen_cost = sum(prob[s] * (data["cdiesel"][t] * value(m.qdiesel[t, s])
                              + data["cPV"][t] * value(m.qPV[t, s]))
                   for t in T for s in S)
    follower_cost = (revenue
                     + sum(prob[s] * data["pPG"][t] * value(m.qPG[t, s])
                           for t in T for s in S)
                     + data["pVOLL"] * qcri)

    return {
        "profit": value(m.profit),
        "revenue": revenue,
        "gen_cost": gen_cost,
        "follower_cost": follower_cost,
        "qMGO": qMGO, "qPG": qPG, "qcri": qcri, "qsupp": qsupp, "D": sD,
        "supp_share": 100.0 * qsupp / sD if sD else 0.0,
        "p_mean": sum(value(m.p[t]) for t in T) / len(T),
    }


def commit(season: str, scenario: str | None, opts: dict) -> dict | None:
    """
    Solve for the price schedule the leader commits under one belief.

    ``scenario`` is the outcome believed CERTAIN; ``None`` gives the
    stochastic commitment (all outcomes at their own probabilities).

    Returns ``{"key", "p", "belief_profit"}`` -- ``belief_profit`` being the
    profit the leader EXPECTS under that belief, not what it realizes -- or
    ``None`` if the solve did not reach an optimal/feasible point.
    """
    data = _instance(season, scenario)
    mpec = BilevelMPEC(data)
    m = mpec.build()

    tc = _solve(mpec, opts)
    key = scenario or STOCHASTIC
    if tc not in _OK:
        print(f"  commitment {key}: not solved ({tc})")
        return None

    return {
        "key": key,
        "p": {t: value(m.p[t]) for t in data["T"]},
        "belief_profit": value(m.profit),
    }


def evaluate(season: str, p_sched: dict[int, float], realization: str,
             opts: dict) -> dict | None:
    """
    Realize ``realization`` with the price held at the committed ``p_sched``.

    The price (and the menu binaries behind it) are first stage and therefore
    frozen; the follower still responds through its full KKT system and the MGO
    still redispatches, so this is the equilibrium under the mismatch, not a
    re-evaluation of the committed dispatch.
    """
    data = _instance(season, realization)
    level_of = _pin_price(data, p_sched)

    mpec = BilevelMPEC(data)
    m = mpec.build()
    for t in data["T"]:
        for k in m.K:
            m.b[t, k].fix(1.0 if k == level_of[t] else 0.0)

    tc = _solve(mpec, opts)
    if tc not in _OK:
        print(f"  realization {realization}: not solved ({tc})")
        return None

    rec = _metrics(m, data)
    rec["termination"] = str(tc)
    return rec


# --------------------------------------------------------------------------
# The matrix
# --------------------------------------------------------------------------
def mismatch_matrix(season: str, opts: dict,
                    stochastic: bool = True) -> dict | None:
    """
    Every (commitment, realization) pair for one season.

    Returns ``{"season", "rows", "cols", "commitments", "cells"}`` where
    ``cells[(i, j)]`` is the per-cell record (with ``regret`` filled in), or
    ``None`` if no commitment solved.
    """
    cols = [sc.key for sc in SCENARIOS]
    beliefs = cols + ([None] if stochastic else [])

    print(f"[{season}] committing one price schedule per belief "
          f"({len(beliefs)}) ...")
    commitments: dict[str, dict] = {}
    for belief in beliefs:
        rec = commit(season, belief, opts)
        if rec is None:
            continue
        commitments[rec["key"]] = rec
        p = rec["p"]
        print(f"  {rec['key']:>12}: p in [{min(p.values()):.3f}, "
              f"{max(p.values()):.3f}]  mean {sum(p.values()) / len(p):.3f}  "
              f"expected profit {rec['belief_profit']:9.2f}")

    if not commitments:
        print(f"[{season}] no commitment solved.")
        return None

    rows = [k for k in cols + [STOCHASTIC] if k in commitments]

    print(f"[{season}] realizing each outcome under each commitment "
          f"({len(rows)} x {len(cols)}) ...")
    cells: dict[tuple[str, str], dict] = {}
    for i in rows:
        for j in cols:
            rec = evaluate(season, commitments[i]["p"], j, opts)
            if rec is None:
                continue
            rec["commitment"], rec["realization"] = i, j
            cells[(i, j)] = rec

    # Regret is measured against the diagonal: the same outcome, priced for.
    foresight = {j: cells[(j, j)]["profit"] for j in cols if (j, j) in cells}
    for (i, j), rec in cells.items():
        rec["regret"] = (foresight[j] - rec["profit"]
                         if j in foresight else None)

    return {"season": season, "rows": rows, "cols": cols,
            "commitments": commitments, "cells": cells}


def _matrix(mm: dict, field: str) -> np.ndarray:
    """The (rows x cols) array of one field, with NaN for cells that failed."""
    out = np.full((len(mm["rows"]), len(mm["cols"])), np.nan)
    for r, i in enumerate(mm["rows"]):
        for c, j in enumerate(mm["cols"]):
            rec = mm["cells"].get((i, j))
            if rec is not None and rec.get(field) is not None:
                out[r, c] = rec[field]
    return out


def expected_profits(mm: dict) -> dict[str, float]:
    """Probability-weighted profit of each commitment row, over the outcomes
    that solved (weights renormalized so a dropped outcome does not read as a
    loss)."""
    out = {}
    for i in mm["rows"]:
        num = wgt = 0.0
        for j in mm["cols"]:
            rec = mm["cells"].get((i, j))
            if rec is None:
                continue
            num += SCENARIO_PROB[j] * rec["profit"]
            wgt += SCENARIO_PROB[j]
        if wgt:
            out[i] = num / wgt
    return out


def perfect_foresight_profit(mm: dict) -> float | None:
    """Expected profit if EDL were known before pricing (the diagonal)."""
    num = wgt = 0.0
    for j in mm["cols"]:
        rec = mm["cells"].get((j, j))
        if rec is None:
            continue
        num += SCENARIO_PROB[j] * rec["profit"]
        wgt += SCENARIO_PROB[j]
    return num / wgt if wgt else None


def summarize(mm: dict) -> list[dict]:
    """One record per commitment: expected profit, worst case, and what the
    belief costs against the stochastic commitment."""
    exp = expected_profits(mm)
    pf = perfect_foresight_profit(mm)
    commitment = exp.get(STOCHASTIC)

    out = []
    for i in mm["rows"]:
        profits = [mm["cells"][(i, j)]["profit"] for j in mm["cols"]
                   if (i, j) in mm["cells"]]
        regrets = [mm["cells"][(i, j)]["regret"] for j in mm["cols"]
                   if (i, j) in mm["cells"]
                   and mm["cells"][(i, j)]["regret"] is not None]
        worst_j = min(((j, mm["cells"][(i, j)]["profit"]) for j in mm["cols"]
                       if (i, j) in mm["cells"]),
                      key=lambda kv: kv[1], default=(None, None))
        out.append({
            "commitment": i,
            "label": label_of(i),
            "prob": SCENARIO_PROB.get(i),
            "belief_profit": mm["commitments"][i]["belief_profit"],
            "exp_profit": exp.get(i),
            "worst_profit": min(profits) if profits else None,
            "worst_case": worst_j[0],
            "max_regret": max(regrets) if regrets else None,
            "exp_regret": (pf - exp[i]) if pf is not None and i in exp else None,
            "cost_vs_commitment": (commitment - exp[i]
                               if commitment is not None and i in exp else None),
        })
    return out


def print_summary(mm: dict) -> None:
    """Console view of the matrix and the headline numbers."""
    season = mm["season"]
    profit = _matrix(mm, "profit")
    exp = expected_profits(mm)
    pf = perfect_foresight_profit(mm)

    width = max(len(label_of(k)) for k in mm["rows"]) + 2
    head = "".join(f"{label_of(j)[:11]:>13}" for j in mm["cols"])
    print(f"\n[{season}] realized leader profit [$]  "
          f"(row = priced for, column = happened)")
    print(" " * width + head + f"{'E[profit]':>13}")
    for r, i in enumerate(mm["rows"]):
        cells = "".join(
            f"{profit[r, c]:>13.2f}" if not np.isnan(profit[r, c]) else f"{'-':>13}"
            for c in range(len(mm["cols"])))
        star = " *" if i == STOCHASTIC else ""
        print(f"{label_of(i):<{width}}{cells}{exp.get(i, float('nan')):>13.2f}{star}")

    if pf is not None:
        print(f"{'Perfect foresight':<{width}}"
              f"{'':>{13 * len(mm['cols'])}}{pf:>13.2f}")
        if STOCHASTIC in exp:
            print(f"  EVPI (perfect foresight - stochastic commitment): {pf - exp[STOCHASTIC]:.2f} $/day")
    if STOCHASTIC in exp:
        worst = min(((i, exp[i]) for i in mm["rows"] if i != STOCHASTIC),
                    key=lambda kv: kv[1], default=None)
        if worst:
            print(f"  costliest wrong belief: {label_of(worst[0])} "
                  f"({exp[STOCHASTIC] - worst[1]:.2f} $/day below the stochastic commitment)")


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def _write_rows(ws, header: list, rows) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "B2"
    for i, label in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(len(str(label)) + 4, 28)


def _sheet_matrix(wb, title: str, mm: dict, field: str, note: str) -> None:
    mat = _matrix(mm, field)
    ws = wb.create_sheet(title[:31])
    _write_rows(
        ws,
        ["priced for \\ happened"] + [label_of(j) for j in mm["cols"]],
        ([label_of(i)] + [None if np.isnan(v) else float(v) for v in mat[r]]
         for r, i in enumerate(mm["rows"])),
    )
    ws.append([])
    ws.append([note])


def save_analysis(results: dict[str, dict], out_path) -> None:
    """Profit / regret / price matrices, the per-cell detail and a summary."""
    wb = Workbook()
    wb.remove(wb.active)

    # ---- summary first: one block per season ----
    ws = wb.create_sheet("summary")
    header = ["season", "commitment", "outcome", "prob (stochastic run)",
              "profit under own belief [$]", "E[profit] over outcomes [$]",
              "worst-case profit [$]", "worst case", "max regret [$]",
              "E[regret] vs perfect foresight [$]", "cost vs stochastic commitment [$]"]
    keys = ["commitment", "label", "prob", "belief_profit", "exp_profit",
            "worst_profit", "worst_case", "max_regret", "exp_regret",
            "cost_vs_commitment"]
    rows = []
    for season, mm in results.items():
        for rec in summarize(mm):
            rows.append([season] + [rec[k] for k in keys])
        pf = perfect_foresight_profit(mm)
        exp = expected_profits(mm)
        rows.append([season, "perfect_foresight", "EDL known before pricing",
                     None, None, pf, None, None, None, 0.0, None])
        if pf is not None and STOCHASTIC in exp:
            rows.append([season, "EVPI",
                         "perfect foresight - stochastic commitment", None, None,
                         pf - exp[STOCHASTIC]])
        rows.append([])
    _write_rows(ws, header, rows)

    for season, mm in results.items():
        _sheet_matrix(
            wb, f"{season}_profit", mm, "profit",
            "Leader profit [$/day] when the row's price schedule meets the "
            "column's EDL outcome. Diagonal = perfect foresight.")
        _sheet_matrix(
            wb, f"{season}_regret", mm, "regret",
            "Regret [$/day] = diagonal profit of the column minus this cell. "
            "0 on the diagonal, >= 0 elsewhere.")
        _sheet_matrix(
            wb, f"{season}_sales", mm, "qMGO",
            "Microgrid sales [kWh/day] realized by each combination.")

        # Committed price schedules, hour by hour.
        ws = wb.create_sheet(f"{season}_prices"[:31])
        hours = sorted(next(iter(mm["commitments"].values()))["p"])
        _write_rows(
            ws, ["priced for"] + [f"t={t}" for t in hours] + ["mean"],
            ([label_of(i)] + [mm["commitments"][i]["p"][t] for t in hours]
             + [sum(mm["commitments"][i]["p"].values()) / len(hours)]
             for i in mm["rows"]),
        )

    # ---- every cell, long form ----
    ws = wb.create_sheet("detail")
    detail_rows = []
    for season, mm in results.items():
        for i in mm["rows"]:
            for j in mm["cols"]:
                rec = mm["cells"].get((i, j))
                if rec is None:
                    continue
                detail_rows.append([season] + [rec.get(k) for k, _ in FIELDS])
    _write_rows(ws, ["season"] + [lbl for _, lbl in FIELDS], detail_rows)

    wb.save(out_path)
    print(f"  wrote {out_path}")


# --------------------------------------------------------------------------
# Plots
# --------------------------------------------------------------------------
# Diverging for profit (loss <- 0 -> gain) and sequential for regret, both
# built from the project palette so the figures read like the rest of the set.
_CMAP_PROFIT = LinearSegmentedColormap.from_list(
    "profit", [COLORS["profit_neg"], "#F7F7F7", COLORS["profit"]])
_CMAP_REGRET = LinearSegmentedColormap.from_list(
    "regret", ["#FFFFFF", COLORS["profit_neg"]])


def _heatmap(ax, mat: np.ndarray, mm: dict, norm, cmap, fmt: str,
             title: str, cbar_label: str) -> None:
    im = ax.imshow(np.ma.masked_invalid(mat), cmap=cmap, norm=norm,
                   aspect="auto")
    cmap.set_bad("#EEEEEE")

    ax.set_xticks(range(len(mm["cols"])),
                  [label_of(j) for j in mm["cols"]], rotation=30, ha="right")
    ax.set_yticks(range(len(mm["rows"])), [label_of(i) for i in mm["rows"]])
    ax.set_xlabel("EDL outcome that happened")
    ax.set_ylabel("EDL outcome the price was set for")
    ax.set_title(title)

    # Cell values, in ink dark or light enough to read on that cell.
    for r in range(mat.shape[0]):
        for c in range(mat.shape[1]):
            v = mat[r, c]
            if np.isnan(v):
                continue
            rgba = cmap(norm(v))
            lum = 0.299 * rgba[0] + 0.587 * rgba[1] + 0.114 * rgba[2]
            ax.text(c, r, format(v, fmt), ha="center", va="center", fontsize=8,
                    color="#FFFFFF" if lum < 0.5 else "#222222")

    # Thin gridlines between the cells (a 2px surface gap, not a heavy border).
    ax.set_xticks(np.arange(-0.5, mat.shape[1], 1), minor=True)
    ax.set_yticks(np.arange(-0.5, mat.shape[0], 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=2)
    ax.tick_params(which="minor", length=0)
    for spine in ax.spines.values():
        spine.set_visible(False)

    cb = ax.figure.colorbar(im, ax=ax, fraction=0.046, pad=0.03)
    cb.set_label(cbar_label, fontsize=8)
    cb.outline.set_visible(False)


def plot_mismatch(mm: dict, out_path) -> None:
    """Profit and regret heatmaps for one season."""
    season = mm["season"]
    profit = _matrix(mm, "profit")
    regret = _matrix(mm, "regret")

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(15, 6))

    lim = np.nanmax(np.abs(profit)) or 1.0
    _heatmap(ax1, profit, mm,
             TwoSlopeNorm(vcenter=0.0, vmin=-lim, vmax=lim), _CMAP_PROFIT,
             ",.1f",
             f"Realized leader profit — {season}\n"
             f"(diagonal = the outcome was priced for)",
             "Profit  [$ / day]")

    rmax = np.nanmax(regret) or 1.0
    _heatmap(ax2, regret, mm, Normalize(vmin=0.0, vmax=rmax), _CMAP_REGRET,
             ",.1f",
             f"Regret vs perfect foresight — {season}\n"
             f"(0 on the diagonal; higher = the wrong belief cost more)",
             "Foregone profit  [$ / day]")

    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)
    print(f"  wrote {out_path}")


def plot_expected(mm: dict, out_path) -> None:
    """Expected profit of each commitment, against the stochastic
    commitment and perfect foresight."""
    season = mm["season"]
    exp = expected_profits(mm)
    pf = perfect_foresight_profit(mm)
    keys = [i for i in mm["rows"] if i in exp]
    vals = [exp[i] for i in keys]

    fig, ax = plt.subplots(figsize=(9, 5.5))
    colors = [COLORS["grid"] if k == STOCHASTIC else
              (COLORS["profit"] if v >= 0 else COLORS["profit_neg"])
              for k, v in zip(keys, vals)]
    ax.bar(range(len(keys)), vals, 0.65, color=colors, edgecolor="white",
           linewidth=2)

    if pf is not None:
        ax.axhline(pf, color=COLORS["shifted"], linewidth=2, linestyle="--")
        ax.text(len(keys) - 0.4, pf, f"  perfect foresight {pf:,.1f} $",
                va="bottom", ha="right", fontsize=8, color=COLORS["shifted"])

    for x, v in enumerate(vals):
        ax.text(x, v, f"{v:,.1f}", ha="center", fontsize=8, color="#222222",
                va="bottom" if v >= 0 else "top")

    ax.set_xticks(range(len(keys)), [label_of(i) for i in keys],
                  rotation=30, ha="right")
    ax.axhline(0, color="#666666", linewidth=0.8)
    ax.set_ylabel("Expected leader profit  [$ / day]")
    ax.set_title(f"Expected profit per price commitment — {season}\n"
                 f"(each schedule met by every EDL outcome at its probability)")
    ax.grid(axis="y", linestyle=":", alpha=0.5)
    ax.margins(y=0.15)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)

    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)
    print(f"  wrote {out_path}")


# --------------------------------------------------------------------------
def main(seasons: list[str] | None = None, stochastic: bool = True) -> None:
    seasons = seasons or ["winter"]
    out_dir = config.OUTPUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)

    opts = {"MIPGap": config.CONFIG["MIPGap"],
            "TimeLimit": config.CONFIG["TimeLimit"]}

    results: dict[str, dict] = {}
    for season in seasons:
        mm = mismatch_matrix(season, opts, stochastic=stochastic)
        if mm is None:
            continue
        results[season] = mm
        print_summary(mm)
        plot_mismatch(mm, out_dir / f"scenario_mismatch_{season}.png")
        plot_expected(mm, out_dir / f"scenario_expected_{season}.png")

    if not results:
        print("No season solved — nothing to write.")
        return

    save_analysis(results, out_dir / "scenario_analysis.xlsx")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Analyse what pricing for the wrong EDL outcome costs.")
    parser.add_argument("--seasons", nargs="+", default=["winter"],
                        choices=list(config.SEASONS.keys()),
                        help="Seasons to analyse (default: winter).")
    parser.add_argument("--no-stochastic", action="store_true",
                        help="Skip the stochastic-commitment row (one solve over "
                             "all outcomes).")
    args = parser.parse_args()
    main(args.seasons, stochastic=not args.no_stochastic)
