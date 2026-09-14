"""
Store a solved model -- every parameter and every variable -- as an Excel file.

A function library only: the run scripts (:mod:`run`, ...) stay the entry
points and call :func:`save_results` after solving.

Works for any Pyomo model in this project (:class:`model.bilevel_mpec.BilevelMPEC`
and any model added later): the components are walked generically, so nothing
has to be listed by hand and new variables show up automatically.

One workbook is written per run. The first and last sheets are fixed:

    scalars             un-indexed parameters, the objective, the run metadata
                        (season, solver, termination) and the config values
                        needed to reproduce the run
    all                 EVERY component in long/tidy form
                        (kind, component, i0, i1, i2, value) -- the complete
                        record

In between there is one pivoted sheet per index signature, one row per index
and one column per component -- ``by_hour`` (T), ``by_hour_scenario`` (T,S),
``by_scenario`` (S), ``by_price_level`` (K), ``by_hour_price_level_scenario``
(T,K,S), and so on. Sheets are grouped by which SETS a component is indexed
over, not by how many indices it has: ``Pi[k]`` and ``prob[s]`` are both
one-dimensional but neither is hourly.

Typical use:

    from output import save_results
    save_results(m, "bilevel_winter", meta={"season": season})
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pyomo.environ import Param, Var, Objective, value

import config


# --------------------------------------------------------------------------
# Collect
# --------------------------------------------------------------------------
def _as_tuple(idx) -> tuple:
    """Normalize a Pyomo index to a tuple: None -> (), scalar -> (scalar,)."""
    if idx is None:
        return ()
    if isinstance(idx, tuple):
        return idx
    return (idx,)


def _sort_key(idx: tuple):
    """Sort mixed int/str indices: numbers first and numerically, then strings."""
    return tuple((0, e, "") if isinstance(e, (int, float)) else (1, 0, str(e))
                 for e in idx)


# Readable names for this project's index sets, used for sheet titles.
_SET_LABEL = {"T": "hour", "S": "scenario", "K": "price_level"}


def _signature(comp) -> tuple[str, ...]:
    """
    The names of the sets a component is indexed over, e.g. ``("T", "S")``.

    Grouping by this rather than by the NUMBER of indices matters: ``Pi[k]``
    and ``prob[s]`` are both one-dimensional but are not hourly, and ``Pi``'s
    integer keys would otherwise collide with the hours in a shared sheet.
    """
    iset = comp.index_set()
    if iset.dimen == 0:
        return ()
    try:
        subs = list(iset.subsets())
    except Exception:
        subs = [iset]
    return tuple(s.name for s in subs)


def collect(model) -> list[dict]:
    """
    Walk ``model`` and return one record per component datum.

    Each record is ``{"kind", "component", "index", "value"}``, where ``kind``
    is ``param`` / ``var`` / ``objective`` and ``index`` is a (possibly empty)
    tuple. Components without a loaded value get ``None`` rather than raising,
    so an unsolved model can still be dumped.
    """
    records: list[dict] = []
    for kind, ctype in (("param", Param), ("var", Var), ("objective", Objective)):
        for comp in model.component_objects(ctype, active=True):
            sig = _signature(comp)
            for idx in comp:
                try:
                    val = value(comp[idx])
                except Exception:
                    val = None          # uninitialized / no value loaded
                records.append({
                    "kind": kind,
                    "component": comp.name,
                    "index": _as_tuple(idx),
                    "signature": sig,
                    "value": val,
                })
    return records


# --------------------------------------------------------------------------
# Sheets
# --------------------------------------------------------------------------
def _write_rows(ws, header: list, rows) -> None:
    """Write a header row in bold plus the data rows, and size the columns."""
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"

    widths = [len(str(h)) for h in header]
    for row in ws.iter_rows(min_row=2, values_only=True):
        for i, v in enumerate(row):
            if i < len(widths):
                widths[i] = max(widths[i], len(str(v)) if v is not None else 0)
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(wdt + 2, 40)


def _sheet_scalars(wb, records: list[dict], meta: dict | None) -> None:
    rows = [(r["kind"], r["component"], r["value"])
            for r in sorted((r for r in records if len(r["index"]) == 0),
                            key=lambda r: (r["kind"], r["component"]))]
    rows += [("meta", k, str(v)) for k, v in (meta or {}).items()]
    # The config is not part of the Pyomo model but is needed to reproduce the
    # run, so it is recorded alongside.
    rows += [("config", k, str(v)) for k, v in config.CONFIG.items()]
    _write_rows(wb.create_sheet("scalars"), ["kind", "name", "value"], rows)


def _sheet_pivot(wb, rows: list[dict], sig: tuple[str, ...]) -> None:
    """One sheet per index signature: one row per index, one column per
    component."""
    if not rows:
        return
    title = ("by_" + "_".join(_SET_LABEL.get(s, s.lower()) for s in sig))[:31]
    index_names = [s.lower() for s in sig]

    components = sorted({r["component"] for r in rows})
    indices = sorted({r["index"] for r in rows}, key=_sort_key)
    table = {(r["index"], r["component"]): r["value"] for r in rows}
    _write_rows(
        wb.create_sheet(title),
        index_names + components,
        (list(idx) + [table.get((idx, c)) for c in components] for idx in indices),
    )


def _sheet_all(wb, records: list[dict]) -> None:
    width = max((len(r["index"]) for r in records), default=0)
    header = ["kind", "component"] + [f"i{k}" for k in range(width)] + ["value"]
    ordered = sorted(records, key=lambda r: (r["kind"], r["component"],
                                             _sort_key(r["index"])))
    rows = ([r["kind"], r["component"],
             *(list(r["index"]) + [None] * (width - len(r["index"]))),
             r["value"]] for r in ordered)
    _write_rows(wb.create_sheet("all"), header, rows)


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------
def save_results(model, prefix: str = "run", out_dir=None,
                 meta: dict | None = None) -> Path:
    """
    Store every parameter and every (solved) variable of ``model`` as one
    Excel workbook.

    Parameters
    ----------
    model : pyomo ConcreteModel
        A built -- normally solved -- model. Components without a value are
        written as blanks rather than raising.
    prefix : str
        Filename stem, e.g. ``"bilevel_winter"`` -> ``bilevel_winter.xlsx``.
    out_dir : path-like, optional
        Target directory (default :data:`config.OUTPUT_DIR`). Created if absent.
    meta : dict, optional
        Extra rows for the ``scalars`` sheet (season, solver, objective,
        termination condition, ...).

    Returns
    -------
    Path
        The workbook written.
    """
    out_dir = Path(out_dir) if out_dir is not None else config.OUTPUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{prefix}.xlsx"

    records = collect(model)

    wb = Workbook()
    wb.remove(wb.active)                      # drop the default empty sheet
    _sheet_scalars(wb, records, meta)

    # One pivoted sheet per index signature (T / T,S / T,K,S / S / K ...),
    # ordered from the simplest index to the most compound.
    groups: dict[tuple[str, ...], list[dict]] = {}
    for r in records:
        if r["index"]:
            groups.setdefault(r["signature"], []).append(r)
    for sig in sorted(groups, key=lambda s: (len(s), s)):
        _sheet_pivot(wb, groups[sig], sig)

    _sheet_all(wb, records)
    wb.save(path)

    print(f"  wrote {path}")
    return path
